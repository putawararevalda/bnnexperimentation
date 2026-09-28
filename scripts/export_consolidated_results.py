"""
Consolidate all training and SEU results into a single Excel file (results_consolidated.xlsx)
with two sheets:
  1. Training_Results
  2. SEU_Results
"""
import csv
import json
import os
import re
from pathlib import Path
import openpyxl

ROOT = Path("E:/bnnexperimentation")
OUTPUT_XLSX = ROOT / "results_consolidated.xlsx"


def safe_float(val):
    if val is None or val == "":
        return None
    try:
        f = float(val)
        return f if (f == f) else None  # handles NaN
    except (ValueError, TypeError):
        return val


def safe_int(val):
    if val is None or val == "":
        return None
    try:
        return int(val)
    except (ValueError, TypeError):
        return val


def collect_training_results():
    rows = []

    # 1. EuroSAT BNN
    eurosat_bnn_root = ROOT / "results/eurosat/bayesian"
    if eurosat_bnn_root.exists():
        for variant_dir in sorted(eurosat_bnn_root.glob("results_eurosat_v02_*")):
            variant_code = variant_dir.name[-2:]
            variant_map = {"00": "base", "01": "smartpool", "02": "dropout", "03": "weight_decay"}
            variant_name = variant_map.get(variant_code, variant_code)

            for cfg_path in sorted(variant_dir.glob("config_*.json")):
                try:
                    with open(cfg_path, "r", encoding="utf-8") as f:
                        cfg = json.load(f)
                except Exception:
                    continue

                ts_match = re.search(r"(\d{8}_\d{6})", cfg_path.name)
                ts = ts_match.group(1) if ts_match else ""

                # read accuracy csv if exists
                best_test_acc = None
                final_test_acc = None
                for acc_path in variant_dir.glob(f"accuracy_results_*{ts}*.csv"):
                    try:
                        with open(acc_path, "r", encoding="utf-8") as f:
                            reader = list(csv.DictReader(f))
                            if reader:
                                final_test_acc = safe_float(reader[-1].get("accuracy"))
                                accs = [safe_float(r.get("accuracy")) for r in reader if safe_float(r.get("accuracy")) is not None]
                                if accs:
                                    best_test_acc = max(accs)
                    except Exception:
                        pass
                    break

                # read losses csv if exists
                final_train_loss = None
                for loss_path in variant_dir.glob(f"losses_*{ts}*.csv"):
                    try:
                        with open(loss_path, "r", encoding="utf-8") as f:
                            reader = list(csv.DictReader(f))
                            if reader:
                                final_train_loss = safe_float(reader[-1].get("loss"))
                    except Exception:
                        pass
                    break

                prior_params = cfg.get("prior_params", {})
                rows.append({
                    "dataset": "EuroSAT",
                    "model_type": "BNN",
                    "fold": 1,
                    "run_context": "paper_clean",
                    "variant": cfg.get("variant", variant_name),
                    "activation": cfg.get("activation"),
                    "prior": cfg.get("prior"),
                    "prior_b": safe_float(prior_params.get("b")),
                    "prior_mu": safe_float(prior_params.get("mu")),
                    "num_epochs": safe_int(cfg.get("num_epochs")),
                    "best_train_acc": safe_float(cfg.get("best_accuracy")),
                    "best_train_acc_epoch": safe_int(cfg.get("best_accuracy_at_epoch")),
                    "final_train_acc": None,
                    "final_train_loss": final_train_loss,
                    "best_val_acc": None,
                    "best_val_acc_epoch": None,
                    "final_val_loss": None,
                    "final_val_acc": None,
                    "best_test_acc": best_test_acc,
                    "final_test_acc": final_test_acc,
                    "batch_size": safe_int(cfg.get("batch_size")),
                    "train_size": safe_int(cfg.get("train_size")),
                    "timestamp": ts,
                    "source_file": str(cfg_path.relative_to(ROOT)),
                })

    # 2. ShipsNet Fold 1 BNN (Paper Baseline)
    shipsnet_bnn_root = ROOT / "results/shipsnet/bayesian"
    if shipsnet_bnn_root.exists():
        for variant_dir in sorted(shipsnet_bnn_root.glob("results_shipsnet_v02_*")):
            variant_code = variant_dir.name[-2:]
            variant_map = {"00": "base", "01": "smartpool", "02": "dropout", "03": "weight_decay"}
            variant_name = variant_map.get(variant_code, variant_code)

            for cfg_path in sorted(variant_dir.glob("config_*.json")):
                try:
                    with open(cfg_path, "r", encoding="utf-8") as f:
                        cfg = json.load(f)
                except Exception:
                    continue

                ts_match = re.search(r"(\d{8}_\d{6})", cfg_path.name)
                ts = ts_match.group(1) if ts_match else ""

                best_test_acc = None
                final_test_acc = None
                for acc_path in variant_dir.glob(f"accuracy_results_*{ts}*.csv"):
                    try:
                        with open(acc_path, "r", encoding="utf-8") as f:
                            reader = list(csv.DictReader(f))
                            if reader:
                                final_test_acc = safe_float(reader[-1].get("accuracy"))
                                accs = [safe_float(r.get("accuracy")) for r in reader if safe_float(r.get("accuracy")) is not None]
                                if accs:
                                    best_test_acc = max(accs)
                    except Exception:
                        pass
                    break

                final_train_loss = None
                for loss_path in variant_dir.glob(f"losses_*{ts}*.csv"):
                    try:
                        with open(loss_path, "r", encoding="utf-8") as f:
                            reader = list(csv.DictReader(f))
                            if reader:
                                final_train_loss = safe_float(reader[-1].get("loss"))
                    except Exception:
                        pass
                    break

                prior_params = cfg.get("prior_params", {})
                rows.append({
                    "dataset": "ShipsNet",
                    "model_type": "BNN",
                    "fold": 1,
                    "run_context": "paper_baseline",
                    "variant": cfg.get("variant", variant_name),
                    "activation": cfg.get("activation"),
                    "prior": cfg.get("prior"),
                    "prior_b": safe_float(prior_params.get("b")),
                    "prior_mu": safe_float(prior_params.get("mu")),
                    "num_epochs": safe_int(cfg.get("num_epochs")),
                    "best_train_acc": safe_float(cfg.get("best_accuracy")),
                    "best_train_acc_epoch": safe_int(cfg.get("best_accuracy_at_epoch")),
                    "final_train_acc": None,
                    "final_train_loss": final_train_loss,
                    "best_val_acc": None,
                    "best_val_acc_epoch": None,
                    "final_val_loss": None,
                    "final_val_acc": None,
                    "best_test_acc": best_test_acc,
                    "final_test_acc": final_test_acc,
                    "batch_size": safe_int(cfg.get("batch_size")),
                    "train_size": safe_int(cfg.get("train_size")),
                    "timestamp": ts,
                    "source_file": str(cfg_path.relative_to(ROOT)),
                })

    # 3. ShipsNet Folds 1–5 BNN (Cross-Validation)
    for fold_num in range(1, 6):
        fold_dir = ROOT / f"results/shipsnet/bayesian/fold{fold_num}"
        if not fold_dir.exists():
            continue
        for variant_subdir in sorted(fold_dir.iterdir()):
            if not variant_subdir.is_dir() or variant_subdir.name.startswith("_"):
                continue
            variant_name = variant_subdir.name

            for cfg_path in sorted(variant_subdir.glob("config_*.json")):
                try:
                    with open(cfg_path, "r", encoding="utf-8") as f:
                        cfg = json.load(f)
                except Exception:
                    continue

                ts_match = re.search(r"(\d{8}_\d{6})", cfg_path.name)
                ts = ts_match.group(1) if ts_match else ""

                best_test_acc = None
                final_test_acc = None
                for acc_path in variant_subdir.glob(f"accuracy_results_*{ts}*.csv"):
                    try:
                        with open(acc_path, "r", encoding="utf-8") as f:
                            reader = list(csv.DictReader(f))
                            if reader:
                                final_test_acc = safe_float(reader[-1].get("accuracy"))
                                accs = [safe_float(r.get("accuracy")) for r in reader if safe_float(r.get("accuracy")) is not None]
                                if accs:
                                    best_test_acc = max(accs)
                    except Exception:
                        pass
                    break

                final_train_loss = None
                for loss_path in variant_subdir.glob(f"losses_*{ts}*.csv"):
                    try:
                        with open(loss_path, "r", encoding="utf-8") as f:
                            reader = list(csv.DictReader(f))
                            if reader:
                                final_train_loss = safe_float(reader[-1].get("loss"))
                    except Exception:
                        pass
                    break

                prior_params = cfg.get("prior_params", {})
                rows.append({
                    "dataset": "ShipsNet",
                    "model_type": "BNN",
                    "fold": fold_num,
                    "run_context": "kfold_cv",
                    "variant": cfg.get("variant", variant_name),
                    "activation": cfg.get("activation"),
                    "prior": cfg.get("prior"),
                    "prior_b": safe_float(prior_params.get("b")),
                    "prior_mu": safe_float(prior_params.get("mu")),
                    "num_epochs": safe_int(cfg.get("num_epochs")),
                    "best_train_acc": safe_float(cfg.get("best_accuracy")),
                    "best_train_acc_epoch": safe_int(cfg.get("best_accuracy_at_epoch")),
                    "final_train_acc": None,
                    "final_train_loss": final_train_loss,
                    "best_val_acc": None,
                    "best_val_acc_epoch": None,
                    "final_val_loss": None,
                    "final_val_acc": None,
                    "best_test_acc": best_test_acc,
                    "final_test_acc": final_test_acc,
                    "batch_size": safe_int(cfg.get("batch_size")),
                    "train_size": safe_int(cfg.get("train_size")),
                    "timestamp": ts,
                    "source_file": str(cfg_path.relative_to(ROOT)),
                })

    # 4. ShipsNet Fold 1 DNN (Paper Baseline)
    det_root = ROOT / "results/shipsnet/deterministic"
    if det_root.exists():
        for variant_dir in sorted(det_root.glob("results_shipsnet_deterministic_0*")):
            if "_SEU" in variant_dir.name:
                continue
            variant_code = variant_dir.name[-2:]
            variant_map = {"00": "base", "01": "smartpool", "02": "dropout", "03": "weight_decay"}
            variant_name = variant_map.get(variant_code, variant_code)

            for log_path in sorted(variant_dir.glob("training_log_*.csv")):
                try:
                    with open(log_path, "r", encoding="utf-8") as f:
                        log_rows = list(csv.DictReader(f))
                except Exception:
                    continue

                if not log_rows:
                    continue

                ts_match = re.search(r"(\d{8}_\d{6})", log_path.name)
                ts = ts_match.group(1) if ts_match else ""

                # parse activation name from filename
                # e.g. training_log_relu_20250807_170542.csv
                parts = log_path.stem.split("_")
                act_name = parts[2] if len(parts) >= 3 else ""

                best_train_acc = max([safe_float(r.get("train_acc")) for r in log_rows if safe_float(r.get("train_acc")) is not None] or [None])
                best_val_acc = max([safe_float(r.get("val_acc")) for r in log_rows if safe_float(r.get("val_acc")) is not None] or [None])

                last_r = log_rows[-1]
                rows.append({
                    "dataset": "ShipsNet",
                    "model_type": "DNN",
                    "fold": 1,
                    "run_context": "paper_baseline",
                    "variant": variant_name,
                    "activation": act_name,
                    "prior": "None",
                    "prior_b": None,
                    "prior_mu": None,
                    "num_epochs": safe_int(last_r.get("epoch")),
                    "best_train_acc": best_train_acc,
                    "best_train_acc_epoch": None,
                    "final_train_acc": safe_float(last_r.get("train_acc")),
                    "final_train_loss": safe_float(last_r.get("train_loss")),
                    "best_val_acc": best_val_acc,
                    "best_val_acc_epoch": None,
                    "final_val_loss": safe_float(last_r.get("val_loss")),
                    "final_val_acc": safe_float(last_r.get("val_acc")),
                    "best_test_acc": None,
                    "final_test_acc": None,
                    "batch_size": 16,
                    "train_size": 3200,
                    "timestamp": ts,
                    "source_file": str(log_path.relative_to(ROOT)),
                })

    # 5. ShipsNet Folds 2–5 DNN
    for fold_num in range(2, 6):
        fold_dir = det_root / f"fold{fold_num}"
        if not fold_dir.exists():
            continue
        for variant_subdir in sorted(fold_dir.iterdir()):
            if not variant_subdir.is_dir():
                continue
            variant_code = variant_subdir.name
            variant_map = {"00": "base", "01": "smartpool", "02": "dropout", "03": "weight_decay"}
            variant_name = variant_map.get(variant_code, variant_code)

            for log_path in sorted(variant_subdir.glob("log_*.csv")):
                try:
                    with open(log_path, "r", encoding="utf-8") as f:
                        log_rows = list(csv.DictReader(f))
                except Exception:
                    continue

                if not log_rows:
                    continue

                ts_match = re.search(r"(\d{8}_\d{6})", log_path.name)
                ts = ts_match.group(1) if ts_match else ""

                parts = log_path.stem.split("_")
                act_name = parts[1] if len(parts) >= 2 else ""

                best_train_acc = max([safe_float(r.get("train_acc")) for r in log_rows if safe_float(r.get("train_acc")) is not None] or [None])
                best_val_acc = max([safe_float(r.get("val_acc")) for r in log_rows if safe_float(r.get("val_acc")) is not None] or [None])

                last_r = log_rows[-1]
                rows.append({
                    "dataset": "ShipsNet",
                    "model_type": "DNN",
                    "fold": fold_num,
                    "run_context": "kfold_cv",
                    "variant": variant_name,
                    "activation": act_name,
                    "prior": "None",
                    "prior_b": None,
                    "prior_mu": None,
                    "num_epochs": safe_int(last_r.get("epoch")),
                    "best_train_acc": best_train_acc,
                    "best_train_acc_epoch": None,
                    "final_train_acc": safe_float(last_r.get("train_acc")),
                    "final_train_loss": safe_float(last_r.get("train_loss")),
                    "best_val_acc": best_val_acc,
                    "best_val_acc_epoch": None,
                    "final_val_loss": safe_float(last_r.get("val_loss")),
                    "final_val_acc": safe_float(last_r.get("val_acc")),
                    "best_test_acc": None,
                    "final_test_acc": None,
                    "batch_size": 16,
                    "train_size": 3200,
                    "timestamp": ts,
                    "source_file": str(log_path.relative_to(ROOT)),
                })

    return rows


def stream_seu_results(ws, seu_headers):
    # EuroSAT BNN (clean)
    eurosat_seu_root = ROOT / "results/eurosat/seu_clean"
    if eurosat_seu_root.exists():
        for variant_dir in sorted(eurosat_seu_root.glob("v02_*")):
            for csv_file in sorted(variant_dir.glob("*.csv")):
                try:
                    with open(csv_file, "r", encoding="utf-8") as f:
                        for row in csv.DictReader(f):
                            ws.append([
                                "EuroSAT",
                                "BNN",
                                1,
                                "paper_clean",
                                row.get("variant"),
                                row.get("activation_fn"),
                                row.get("prior"),
                                safe_float(row.get("prior_b")),
                                safe_float(row.get("prior_mu")),
                                safe_float(row.get("best_accuracy")),
                                safe_float(row.get("initial_accuracy")),
                                row.get("location_layer"),
                                row.get("location_module"),
                                safe_int(row.get("location_index")),
                                safe_int(row.get("bit_index")),
                                row.get("param_type"),
                                safe_float(row.get("accuracy_after_seu")),
                                safe_float(row.get("accuracy_change")),
                                safe_float(row.get("softmax_difference")),
                                safe_float(row.get("mean_abs_difference")),
                                safe_int(row.get("original_bit_condition")),
                                row.get("remarks", ""),
                                str(csv_file.relative_to(ROOT)),
                            ])
                except Exception as e:
                    print(f"Error reading {csv_file}: {e}")

    # ShipsNet Fold 1 BNN (Paper Baseline)
    shipsnet_seu_root = ROOT / "results/shipsnet/seu"
    if shipsnet_seu_root.exists():
        for variant_dir in sorted(shipsnet_seu_root.glob("results_shipsnet_v02_*_SEU")):
            if "median" in variant_dir.name:
                continue
            for csv_file in sorted(variant_dir.glob("*.csv")):
                try:
                    with open(csv_file, "r", encoding="utf-8") as f:
                        for row in csv.DictReader(f):
                            ws.append([
                                "ShipsNet",
                                "BNN",
                                1,
                                "paper_baseline",
                                row.get("variant"),
                                row.get("activation_fn"),
                                row.get("prior"),
                                safe_float(row.get("prior_b")),
                                safe_float(row.get("prior_mu")),
                                safe_float(row.get("best_accuracy")),
                                safe_float(row.get("initial_accuracy")),
                                row.get("location_layer"),
                                row.get("location_module"),
                                safe_int(row.get("location_index")),
                                safe_int(row.get("bit_index")),
                                row.get("param_type"),
                                safe_float(row.get("accuracy_after_seu")),
                                safe_float(row.get("accuracy_change")),
                                safe_float(row.get("softmax_difference")),
                                safe_float(row.get("mean_abs_difference")),
                                safe_int(row.get("original_bit_condition")),
                                row.get("remarks", ""),
                                str(csv_file.relative_to(ROOT)),
                            ])
                except Exception as e:
                    print(f"Error reading {csv_file}: {e}")

    # ShipsNet Folds 1–5 BNN
    for fold_num in range(1, 6):
        fold_seu_dir = ROOT / f"results/shipsnet/seu/fold{fold_num}"
        if not fold_seu_dir.exists():
            continue
        for variant_dir in sorted(fold_seu_dir.iterdir()):
            if not variant_dir.is_dir():
                continue
            for csv_file in sorted(variant_dir.glob("*.csv")):
                try:
                    with open(csv_file, "r", encoding="utf-8") as f:
                        for row in csv.DictReader(f):
                            ws.append([
                                "ShipsNet",
                                "BNN",
                                safe_int(row.get("fold", fold_num)),
                                "kfold_cv",
                                row.get("variant"),
                                row.get("activation_fn"),
                                row.get("prior"),
                                safe_float(row.get("prior_b")),
                                safe_float(row.get("prior_mu")),
                                safe_float(row.get("best_accuracy")),
                                safe_float(row.get("initial_accuracy")),
                                row.get("location_layer"),
                                row.get("location_module"),
                                safe_int(row.get("location_index")),
                                safe_int(row.get("bit_index")),
                                row.get("param_type"),
                                safe_float(row.get("accuracy_after_seu")),
                                safe_float(row.get("accuracy_change")),
                                safe_float(row.get("softmax_difference")),
                                safe_float(row.get("mean_abs_difference")),
                                safe_int(row.get("original_bit_condition")),
                                row.get("remarks", ""),
                                str(csv_file.relative_to(ROOT)),
                            ])
                except Exception as e:
                    print(f"Error reading {csv_file}: {e}")

    # ShipsNet Fold 1 DNN (Paper Baseline)
    det_root = ROOT / "results/shipsnet/deterministic"
    if det_root.exists():
        for variant_dir in sorted(det_root.glob("results_shipsnet_deterministic_0*_SEU")):
            for csv_file in sorted(variant_dir.glob("*.csv")):
                try:
                    with open(csv_file, "r", encoding="utf-8") as f:
                        for row in csv.DictReader(f):
                            variant_code = row.get("model_variant", "")
                            variant_map = {"00": "base", "01": "smartpool", "02": "dropout", "03": "weight_decay"}
                            variant_name = variant_map.get(variant_code, variant_code)
                            ws.append([
                                "ShipsNet",
                                "DNN",
                                safe_int(row.get("fold", 1)),
                                "paper_baseline",
                                variant_name,
                                row.get("activation_fn"),
                                "None",
                                None,
                                None,
                                None,
                                safe_float(row.get("initial_accuracy")),
                                row.get("location_layer"),
                                row.get("location_module"),
                                safe_int(row.get("location_index")),
                                safe_int(row.get("bit_index")),
                                "weight_or_bias",
                                safe_float(row.get("accuracy_after_seu")),
                                safe_float(row.get("accuracy_change")),
                                safe_float(row.get("softmax_difference")),
                                safe_float(row.get("mean_abs_difference")),
                                safe_int(row.get("original_bit_condition")),
                                row.get("remarks", ""),
                                str(csv_file.relative_to(ROOT)),
                            ])
                except Exception as e:
                    print(f"Error reading {csv_file}: {e}")


def main():
    print("Collecting Training Results...")
    training_rows = collect_training_results()
    print(f"Collected {len(training_rows)} training run records.")

    print("Creating Excel workbook...")
    wb = openpyxl.Workbook(write_only=True)

    # 1. Training_Results Sheet
    ws_train = wb.create_sheet(title="Training_Results")
    train_headers = [
        "dataset", "model_type", "fold", "run_context", "variant", "activation",
        "prior", "prior_b", "prior_mu", "num_epochs", "best_train_acc",
        "best_train_acc_epoch", "final_train_acc", "final_train_loss",
        "best_val_acc", "best_val_acc_epoch", "final_val_loss", "final_val_acc",
        "best_test_acc", "final_test_acc", "batch_size", "train_size",
        "timestamp", "source_file"
    ]
    ws_train.append(train_headers)
    for r in training_rows:
        ws_train.append([r.get(h) for h in train_headers])

    # 2. SEU_Results Sheet
    ws_seu = wb.create_sheet(title="SEU_Results")
    seu_headers = [
        "dataset", "model_type", "fold", "run_context", "variant", "activation_fn",
        "prior", "prior_b", "prior_mu", "best_accuracy", "initial_accuracy",
        "location_layer", "location_module", "location_index", "bit_index",
        "param_type", "accuracy_after_seu", "accuracy_change", "softmax_difference",
        "mean_abs_difference", "original_bit_condition", "remarks", "source_file"
    ]
    ws_seu.append(seu_headers)

    print("Streaming SEU Results into worksheet...")
    stream_seu_results(ws_seu, seu_headers)

    print(f"Saving workbook to {OUTPUT_XLSX}...")
    wb.save(OUTPUT_XLSX)
    print("Done! Consolidated file created successfully.")


if __name__ == "__main__":
    main()
